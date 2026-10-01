"""검토 도구: 오프라인 review.html 생성, 검토 결과를 합친 responses_final.xlsx 내보내기.

review.html은 output 폴더 안에 만들어지며 크롭 이미지를 상대경로로 참조한다.
검토 결과는 브라우저에서 review_decisions.csv로 내려받아 output 폴더에 둔다.
"""

import csv
import json
from collections import Counter
from pathlib import Path

from .kits import load_kits, normalize_kit

HW_FIELDS = ["grade", "kit_2", "kit_3", "q5"]
HW_LABELS = {"grade": "학년", "kit_2": "2교시 키트", "kit_3": "3교시 키트", "q5": "5. 하고 싶은 말"}
INVALID = "무효"
DECISION_COLS = ["id", "field", "auto_value", "final_value"]


def read_rows(out_dir: Path) -> list[dict]:
    with (out_dir / "responses.csv").open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def needs_review(row: dict) -> bool:
    return row["omr_status"] == "needs_review" or any(row.get(f"{f}_conf") == "low" for f in HW_FIELDS)


def build_review(out_dir: Path, questions: list[dict]) -> Path:
    items = []
    prior = {i for i, _ in read_decisions(out_dir / "review_decisions.csv")}  # 이미 검토 기록이 있는 ID
    for r in read_rows(out_dir):
        fields = []
        for q in questions:
            flags = [f.split(":", 1)[1] for f in r["omr_flags"].split(";") if f.startswith(q["id"] + ":")]
            fields.append({"id": q["id"], "label": q["text"], "kind": "choice", "value": r[q["id"]],
                           "candidate": r[q["id"] + "_candidate"], "flag": ",".join(flags),
                           "options": [o["label"] for o in q["options"]] + [INVALID]})
        for f in HW_FIELDS:
            blank = r[f"blank_{f}"] == "True"
            flag = "low" if r[f"{f}_conf"] == "low" else ""
            if blank:  # 빈 칸으로 판정돼도 다른 칸에 쓴 내용을 옮겨 입력할 수 있게 입력칸은 남긴다.
                flag = flag or ""
            fields.append({"id": f, "label": HW_LABELS[f], "kind": "text", "value": r[f], "flag": flag,
                           "blank": blank, "crop": "" if blank else r[f"crop_{f}"]})
        debug = Path("debug") / f"{r['id']}.jpg"
        orig = Path("originals") / f"{r['id']}.jpg"
        items.append({"id": r["id"], "school": r["school"], "file": r["file"],
                      "original": orig.as_posix() if (out_dir / orig).exists() else "", "debug": debug.as_posix() if (out_dir / debug).exists() else "",
                      "omr_review": r["omr_status"] == "needs_review", "needs": needs_review(r), "prior": r["id"] in prior,
                      "q5": bool(r["q5"]), "fields": fields})
    path = out_dir / "review.html"
    path.write_text(HTML.replace("__DATA__", json.dumps({"items": items, "invalid": INVALID}, ensure_ascii=False)
                                 .replace("</", "<\\/")), encoding="utf-8")
    return path


def read_decisions(path: Path) -> dict[tuple[str, str], str]:
    if not path.exists():
        return {}
    with path.open(newline="", encoding="utf-8-sig") as f:
        return {(d["id"], d["field"]): d["final_value"] for d in csv.DictReader(f)}


def read_links(path: Path) -> dict[str, str]:
    """drive_links.csv(file,url) → {키: URL}. file은 상대경로 또는 파일명 모두 허용한다."""
    with path.open(newline="", encoding="utf-8-sig") as f:
        return {d["file"].strip(): d["url"].strip() for d in csv.DictReader(f) if d.get("file") and d.get("url")}


def find_link(links: dict[str, str], file: str) -> str:
    """상대경로 → 파일명 → 확장자 뺀 이름 순으로 찾는다(드라이브에 jpg 사본을 올려도 매칭되도록)."""
    name = file.rsplit("/", 1)[-1]
    stem = name.rsplit(".", 1)[0]
    by_stem = {k.rsplit("/", 1)[-1].rsplit(".", 1)[0]: v for k, v in links.items()}
    return links.get(file) or links.get(name) or by_stem.get(stem, "")


def read_corrections(path: Path) -> dict[tuple[str, str], str]:
    """corrections.csv(id,field,value): 검토 결과 위에 덮어쓰는 수동 보정. 빈 value는 '비움'을 뜻한다."""
    if not path.exists():
        return {}
    with path.open(newline="", encoding="utf-8-sig") as f:
        return {(d["id"], d["field"]): d["value"] for d in csv.DictReader(f)}


def export_xlsx(out_dir: Path, questions: list[dict], decisions_path: Path, links: dict[str, str] | None = None) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = read_rows(out_dir)
    dec = read_decisions(decisions_path)
    corr = read_corrections(out_dir / "corrections.csv")
    reviewed_ids = {i for i, _ in dec}
    qids = [q["id"] for q in questions]
    kits = load_kits()
    final, pending, log, kit_issues = [], [], [], []
    for r in rows:
        out = {"id": r["id"], "school": r["school"], "file": r["file"],
               "share_url": find_link(links, r["file"]) if links else ""}
        for f in qids + HW_FIELDS:
            auto = r[f]
            out[f] = dec.get((r["id"], f), auto)
            if (r["id"], f) in dec:
                log.append([r["id"], f, auto, out[f], "수정" if out[f] != auto else "확인"])
            if (r["id"], f) in corr:
                log.append([r["id"], f, out[f], corr[(r["id"], f)], "보정"])
                out[f] = corr[(r["id"], f)]
        for k in ("kit_2", "kit_3"):
            out[k + "_std"], state = normalize_kit(out[k], kits)
            if state in ("unmatched", "ambiguous"):
                kit_issues.append((r["id"], k, out[k]))
        out["kit_check"] = "확인필요" if any(i[0] == r["id"] for i in kit_issues) else ""
        if needs_review(r) and r["id"] not in reviewed_ids:
            out["status"] = "needs_review"
            pending.append(r["id"])
        else:
            out["status"] = "reviewed" if r["id"] in reviewed_ids else "auto"
        final.append(out)

    wb = Workbook()
    ws = wb.active
    ws.title = "응답"
    cols = ["id", "school", "file", "share_url"] + qids + ["grade", "kit_2", "kit_2_std", "kit_3", "kit_3_std", "kit_check", "q5", "status"]
    ws.append(cols)
    link_col = cols.index("share_url") + 1
    for o in final:
        ws.append([o[c] for c in cols])
        if o["share_url"]:
            cell = ws.cell(row=ws.max_row, column=link_col)
            cell.hyperlink = o["share_url"]
            cell.style = "Hyperlink"
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"

    agg = wb.create_sheet("집계")
    schools = sorted({o["school"] for o in final})
    for q in questions:
        agg.append([f"{q['id']}. {q['text']}"])
        agg[agg.max_row][0].font = Font(bold=True)
        labels = [o["label"] for o in q["options"]] + [INVALID, ""]
        agg.append(["선택지", "전체"] + schools)
        for lab in labels:
            counts = Counter((o["school"], o[q["id"]]) for o in final)
            agg.append([lab or "(무응답)", sum(counts[(s, lab)] for s in schools)] + [counts[(s, lab)] for s in schools])
        agg.append([])
    invalid = [(o["id"], q, o["file"], o["share_url"]) for o in final for q in qids if o[q] == INVALID]
    if invalid:  # 무효 판정은 담당자에게 검토를 요청한다
        req = wb.create_sheet("검토요청")
        req.append(["id", "문항", "원본 파일", "공유링크", "사유"])
        for row in invalid:
            req.append(list(row) + ["무효 판정(두 칸 이상 의도적 표시)"])
    kit_ws = wb.create_sheet("키트")
    kit_ws.append(["표준 키트명", "2교시", "3교시", "합계"] + schools)
    for c in kit_ws[1]:
        c.font = Font(bold=True)
    def count(name: str, slot: str, school: str | None = None) -> int:
        # 무응답("")은 표준명이 없고 원문도 비어 있는 경우만 센다(미분류와 구분).
        return sum(1 for o in final if (school is None or o["school"] == school)
                   and o[slot + "_std"] == name and (name or o[slot] == ""))

    for name in list(kits) + [""]:
        c2, c3 = count(name, "kit_2"), count(name, "kit_3")
        kit_ws.append([name or "(무응답)", c2, c3, c2 + c3]
                      + [count(name, "kit_2", sc) + count(name, "kit_3", sc) for sc in schools])
    if kit_issues:
        kit_ws.append([])
        kit_ws.append(["표준명을 정하지 못한 값(검토 화면에서 고친 뒤 다시 export)"])
        kit_ws.append(["id", "칸", "값"])
        for row in kit_issues:
            kit_ws.append(list(row))
    hist = wb.create_sheet("검토이력")
    hist.append(["id", "field", "자동값", "최종값", "구분"])
    for row in log:
        hist.append(row)
    path = out_dir / "responses_final.xlsx"
    wb.save(path)
    with (out_dir / "responses_final.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(final)
    return {"path": path, "rows": len(final), "pending": pending, "edited": sum(1 for r in log if r[4] == "수정"),
            "no_link": [o["file"] for o in final if links is not None and not o["share_url"]],
            "kit_issues": kit_issues, "invalid": invalid}


HTML = """<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>설문 검토</title>
<style>
:root{--bg:#f6f7f9;--fg:#1d2330;--card:#fff;--line:#d9dde5;--accent:#2f6fed;--warn:#fff3cd;--warnline:#e0a800;--ok:#d7f0dd}
@media (prefers-color-scheme:dark){:root{--bg:#14171c;--fg:#e6e9ef;--card:#1d222b;--line:#333a47;--accent:#6c9cff;--warn:#4a3d10;--warnline:#c79a1a;--ok:#1f3d29}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.5 system-ui,"Malgun Gothic",sans-serif}
header{position:sticky;top:0;z-index:5;background:var(--card);border-bottom:1px solid var(--line);padding:10px 16px;display:flex;gap:12px;flex-wrap:wrap;align-items:center}
header b{font-size:16px}button,select,input,textarea{font:inherit;color:inherit;background:var(--card);border:1px solid var(--line);border-radius:6px;padding:5px 9px}
button{cursor:pointer}button.primary{background:var(--accent);color:#fff;border-color:var(--accent)}
main{max-width:1000px;margin:0 auto;padding:16px}.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px;margin-bottom:14px}
.card.done{background:var(--ok)}.card.cur{outline:2px solid var(--accent)}.head{display:flex;justify-content:space-between;gap:10px;flex-wrap:wrap;align-items:center;margin-bottom:8px}.head .hint{flex:1;user-select:all}
.field{display:grid;grid-template-columns:150px 1fr;gap:8px;padding:6px 0;border-top:1px solid var(--line)}.field.flag{background:var(--warn);border-left:4px solid var(--warnline);padding-left:8px}
.field img{max-width:100%;max-height:220px;border:1px solid var(--line);background:#fff;display:block;margin-bottom:6px}.lab{font-size:13px;opacity:.8}
textarea{width:100%;min-height:54px}input[type=text]{width:100%}.dbg{max-width:100%;border:1px solid var(--line);margin:6px 0}
.hint{font-size:12px;opacity:.7}input[type=search]{min-width:160px}@media (max-width:640px){.field{grid-template-columns:1fr}}
</style></head><body>
<header><b>설문 검토</b><span id="prog"></span>
<select id="filter"><option value="todo">미검토만</option><option value="omr">체크박스 플래그</option><option value="q5">자유의견 미검토</option><option value="all">전체</option></select>
<input id="q" type="search" placeholder="ID·파일명 검색" size="22">
<button id="csv" class="primary">결과 CSV 다운로드</button>
<span class="hint">j/k 이동 · Enter 확인하고 다음 · Esc 입력 해제</span></header>
<main id="list"></main>
<script>
const DATA=__DATA__;const KEY="survey-review:"+location.pathname;
let st={};try{st=JSON.parse(localStorage.getItem(KEY)||"{}")}catch(e){}
const save=()=>{try{localStorage.setItem(KEY,JSON.stringify(st))}catch(e){}};
const $=s=>document.querySelector(s);let cur=0,shown=[];
const val=(it,f)=>(st[it.id]&&st[it.id].v&&f.id in st[it.id].v)?st[it.id].v[f.id]:f.value;
function visible(){const m=$("#filter").value,q=$("#q").value.trim().toLowerCase();return DATA.items.filter(it=>(!q||(it.id+it.file).toLowerCase().includes(q))&&(q||m==="all"||(m==="omr"?it.omr_review:m==="q5"?it.q5&&!it.prior&&!(st[it.id]&&st[it.id].done):it.needs&&!(st[it.id]&&st[it.id].done))))}
function render(){shown=visible();if(cur>=shown.length)cur=Math.max(0,shown.length-1);const done=DATA.items.filter(i=>st[i.id]&&st[i.id].done).length;
$("#prog").textContent=`${done}/${DATA.items.length} 검토 완료`;const box=$("#list");box.innerHTML="";
shown.forEach((it,idx)=>{const c=document.createElement("div");c.className="card"+(st[it.id]&&st[it.id].done?" done":"")+(idx===cur?" cur":"");c.id="c"+idx;
c.innerHTML=`<div class="head"><b>#${it.id} · ${it.school}</b><span class="hint">${it.file}</span>${it.original?`<a href="${it.original}" target="_blank">원본 보기</a>`:""}<button data-ok="${idx}">확인</button></div>`;
if(it.debug&&it.omr_review){const im=document.createElement("img");im.src=it.debug;im.className="dbg";im.loading="lazy";c.appendChild(im)}
it.fields.forEach(f=>{const row=document.createElement("div");row.className="field"+(f.flag?" flag":"");const l=document.createElement("div");l.innerHTML=`<div>${f.label}</div><div class="lab">${f.flag?"플래그: "+f.flag+(f.candidate?" (후보: "+f.candidate+")":""):""}</div>`;
const r=document.createElement("div");if(f.blank){const n=document.createElement("div");n.className="hint";n.textContent="빈 칸으로 판정됨 — 다른 칸에 쓴 내용이 있으면 여기에 입력";r.appendChild(n)}if(f.crop){const im=document.createElement("img");im.src=f.crop;im.loading="lazy";r.appendChild(im)}
let el;if(f.kind==="choice"){el=document.createElement("select");[""].concat(f.options).forEach(o=>{const op=document.createElement("option");op.value=o;op.textContent=o||"(무응답)";el.appendChild(op)});}
else{el=f.id==="q5"?document.createElement("textarea"):Object.assign(document.createElement("input"),{type:"text"})}
el.value=val(it,f);el.dataset.f=f.id;el.addEventListener("input",()=>{(st[it.id]=st[it.id]||{}).v=st[it.id].v||{};st[it.id].v[f.id]=el.value;save()});
r.appendChild(el);row.append(l,r);c.appendChild(row)});box.appendChild(c)});
document.querySelectorAll("[data-ok]").forEach(b=>b.onclick=()=>confirmAt(+b.dataset.ok))}
function confirmAt(i){const it=shown[i];if(!it)return;(st[it.id]=st[it.id]||{}).done=true;save();const next=i;render();cur=Math.min(next,shown.length-1);focus()}
function focus(){const c=$("#c"+cur);if(c){c.scrollIntoView({block:"start",behavior:"smooth"});document.querySelectorAll(".card").forEach(x=>x.classList.remove("cur"));c.classList.add("cur")}}
document.addEventListener("keydown",e=>{const t=e.target.tagName;if(e.key==="Escape"){document.activeElement.blur();return}
if(e.key==="Enter"&&(t!=="TEXTAREA"||e.ctrlKey)){e.preventDefault();confirmAt(cur);return}
if(t==="INPUT"||t==="TEXTAREA"||t==="SELECT")return;if(e.key==="j"){cur=Math.min(cur+1,shown.length-1);focus()}if(e.key==="k"){cur=Math.max(cur-1,0);focus()}});
$("#filter").onchange=()=>{cur=0;render()};$("#q").oninput=()=>{cur=0;render()};
$("#csv").onclick=()=>{const q=s=>'"'+String(s).replace(/"/g,'""')+'"';const lines=["id,field,auto_value,final_value"];
DATA.items.forEach(it=>{if(!(st[it.id]&&st[it.id].done))return;it.fields.forEach(f=>lines.push([it.id,f.id,f.value,val(it,f)].map(q).join(",")))});
const a=document.createElement("a");a.href=URL.createObjectURL(new Blob(["\\ufeff"+lines.join("\\r\\n")],{type:"text/csv;charset=utf-8"}));a.download="review_decisions.csv";a.click()};
render();
</script></body></html>
"""
